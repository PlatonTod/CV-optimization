"""
Dimax HR v2 — ядро: вакансії з критеріями, промпти, імпорт CSV з Betterteam,
читання резюме (посилання / файли / ZIP), оцінка ШІ за кількома вакансіями, звіт Excel.
"""
import concurrent.futures
import hashlib
import io
import json
import os
import re
import threading
import time
import unicodedata
import zipfile
from dataclasses import dataclass
from datetime import datetime
from typing import Callable, Literal, Optional

import pandas as pd
import requests
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

GEMINI_URL = "https://generativelanguage.googleapis.com/v1beta/openai/"
DEFAULT_MODEL = {"gemini": "gemini-3.8-flash", "openai": "gpt-4o-mini"}
MODES = {"paid": {"workers": 8, "pause": 0.0}, "free": {"workers": 1, "pause": 5.0}}
MAX_CV_CHARS = 7000

LEVELS = ["базовий", "середній", "вільний", "рідний"]
REQ = {"must": "обов'язково", "pref": "бажано"}
LOC_MODES = {"must": "кандидат має зараз бути там (обов'язково)", "pref": "бажано там",
             "reloc": "будь-де, якщо готовий переїхати", "any": "не важливо"}
EXP_OPTIONS = ["", "Без досвіду / junior", "1–2 роки", "2–5 років", "5–10 років", "Понад 10 років"]
EDU_OPTIONS = ["", "Будь-яка вища", "Технічна", "Гуманітарна", "Економічна / бізнес"]


# =====================================================================
# Налаштування
# =====================================================================
@dataclass
class Settings:
    api_key: str
    mode: str = "paid"
    provider: str = "gemini"
    model: str = ""
    pause: Optional[float] = None
    workers: Optional[int] = None
    smart_filter: bool = True  # не надсилати в ШІ тих, у кого немає ознак UA/СНД, для вакансій з обов'язковою укр/рос

    @property
    def model_name(self):
        return self.model.strip() or DEFAULT_MODEL[self.provider]

    @property
    def n_workers(self):
        return self.workers or MODES[self.mode]["workers"]

    @property
    def request_pause(self):
        return MODES[self.mode]["pause"] if self.pause is None else self.pause


def stable_hash(s):
    return hashlib.md5(s.encode("utf-8")).hexdigest()[:10]


# =====================================================================
# Вакансії
# =====================================================================
def blank_vacancy():
    return {"id": "", "title": "", "languages": [{"lang": "англійська", "level": "вільний", "req": "must"}],
            "tools": [], "location": {"place": "", "mode": "pref"},
            "experience": {"years": "", "field": "", "req": "pref"}, "education": {"type": "", "req": "pref"},
            "notes": "", "prompt": "", "promptFor": "", "promptEdited": False}


def criteria_text(v):
    out = ["Посада: " + (v.get("title") or "—")]
    langs = [l for l in v.get("languages", []) if str(l.get("lang", "")).strip()]
    if langs:
        out.append("Мови: " + "; ".join(f"{l['lang'].strip()} — {l.get('level', 'вільний')} ({REQ.get(l.get('req'), 'бажано')})" for l in langs))
    tools = [t for t in v.get("tools", []) if str(t.get("name", "")).strip()]
    if tools:
        out.append("Інструменти та навички: " + "; ".join(f"{t['name'].strip()} ({REQ.get(t.get('req'), 'бажано')})" for t in tools))
    loc = v.get("location", {})
    if loc.get("mode") == "any":
        out.append("Локація: не важлива")
    elif str(loc.get("place", "")).strip():
        out.append(f"Локація: {loc['place'].strip()} — {LOC_MODES.get(loc.get('mode'), '')}")
    exp = v.get("experience", {})
    if exp.get("years") or str(exp.get("field", "")).strip():
        field = f" у сфері «{exp['field'].strip()}»" if str(exp.get("field", "")).strip() else ""
        out.append(f"Досвід: {exp.get('years') or 'стаж не важливий'}{field} ({REQ.get(exp.get('req'), 'бажано')})")
    edu = v.get("education", {})
    if edu.get("type"):
        out.append(f"Освіта: {edu['type'].lower()} ({REQ.get(edu.get('req'), 'бажано')})")
    if str(v.get("notes", "")).strip():
        out.append("Інші вимоги: " + v["notes"].strip())
    return "\n".join(out)


def crit_hash(v):
    return stable_hash(criteria_text(v))


def prompt_state(v):
    if not str(v.get("prompt", "")).strip():
        return "neutral", "Промпту ще немає"
    if v.get("promptFor") and v["promptFor"] != crit_hash(v):
        return "warn", "Критерії змінено після генерації"
    return "good", "Промпт відредаговано вручну" if v.get("promptEdited") else "Промпт відповідає критеріям"


def requires_uk_ru(v):
    """Чи є українська/російська обов'язковою мовою (тоді працює економний фільтр)."""
    return any(l.get("req") == "must" and re.search(r"україн|росій|russian|ukrain", str(l.get("lang", "")), re.I)
               for l in v.get("languages", []))


CIS_NOTE = ("Кандидати з України, СНД (Азербайджан, Грузія, Казахстан, Вірменія, Узбекистан, Молдова, Киргизстан, "
            "Таджикистан, Туркменістан, Білорусь) і Балтії часто не вказують українську чи російську в англомовному резюме: "
            "визначай походження за освітою, місцями роботи, містами, телефонним кодом, ім'ям і вважай, що мовою володіють "
            "(перевіримо на дзвінку). Кандидати, які зараз перебувають у Росії, не розглядаються.")
CIS_RULES = ("- Кандидати з України, СНД (Азербайджан, Грузія, Казахстан, Вірменія, Узбекистан, Молдова, Киргизстан, Таджикистан, "
             "Туркменістан, Білорусь) і Балтії часто не вказують українську чи російську в англомовному резюме. Визначай походження "
             "за освітою, місцями роботи, містами, телефонним кодом, ім'ям. Якщо таке походження є — вважай вимогу про укр/рос "
             "виконаною і додай у «бракує»: «перевірити укр/рос на дзвінку».\n"
             "- Кандидат, який зараз перебуває в Росії, отримує 0 балів.")

TEMPLATES = {
    "MEA Sales Manager": {
        "title": "MEA Sales Manager",
        "languages": [{"lang": "арабська", "level": "вільний", "req": "must"}, {"lang": "англійська", "level": "вільний", "req": "must"},
                      {"lang": "українська або російська", "level": "вільний", "req": "must"}],
        "tools": [{"name": "активні B2B-продажі", "req": "must"}, {"name": "пошук і ведення дистриб'юторів / key accounts", "req": "must"},
                  {"name": "продажі технічного обладнання (генератори, енергетика, інструменти)", "req": "pref"}, {"name": "CRM", "req": "pref"}],
        "location": {"place": "Дубай, ОАЕ", "mode": "pref"},
        "experience": {"years": "2–5 років", "field": "B2B-продажі на Близькому Сході та в Африці (MEA)", "req": "pref"},
        "education": {"type": "", "req": "pref"},
        "notes": "Компанія: Dimax Group, генератори та силове обладнання Könner & Söhnen. Особисті переговори з великими клієнтами. " + CIS_NOTE,
        "prompt": "ПОСАДА: MEA Sales Manager у Dimax Group (генератори та силове обладнання Könner & Söhnen). Активні B2B-продажі на Близькому Сході та в Африці: пошук великих дистриб'юторів і ключових клієнтів, особисті переговори.\n\n"
                  "ОБОВ'ЯЗКОВІ ВИМОГИ (якщо хоч одна не виконана — бал не вище 49):\n1. Арабська — вільно або рідна.\n2. Англійська — вільно.\n3. Українська або російська — вільно (див. правило про походження).\n4. Досвід активних B2B-продажів: не роздріб і не обслуговування клієнтів.\n5. Робота з дистриб'юторами або key accounts.\n\n"
                  "БАЖАНІ ВИМОГИ (додають бали):\n- Продажі технічного обладнання: генератори, енергетика, інструменти, промислові товари.\n- Робота з CRM.\n- 2–5 років і більше B2B-продажів саме в регіоні MEA.\n- Проживання в Дубаї або іншому еміраті ОАЕ.\n\n"
                  "ВАГИ (разом 100): мови — 30; B2B-продажі та дистрибуція в MEA — 35; технічний продукт — 15; стаж — 10; локація ОАЕ — 10.\n\n"
                  "ОСОБЛИВІ ПРАВИЛА:\n" + CIS_RULES + "\n\n"
                  "ШКАЛА: 85–100 — ідеальний збіг (усі обов'язкові й більшість бажаних); 70–84 — сильний кандидат; 50–69 — потенційний, є прогалини в бажаних вимогах; 0–49 — не виконано обов'язкові вимоги.",
    },
    "Country Manager (Dubai)": {
        "title": "Country Manager (Dubai)",
        "languages": [{"lang": "англійська", "level": "вільний", "req": "must"}, {"lang": "українська або російська", "level": "вільний", "req": "must"},
                      {"lang": "арабська", "level": "базовий", "req": "pref"}],
        "tools": [{"name": "відкриття та організація офісу з нуля", "req": "must"},
                  {"name": "робота з юристами, аутсорс- і аутстаф-компаніями, договорами", "req": "must"},
                  {"name": "адміністративне та операційне управління", "req": "must"},
                  {"name": "реєстрація компанії, візи, ліцензії в ОАЕ", "req": "pref"}],
        "location": {"place": "Дубай або сусідні емірати, ОАЕ", "mode": "must"},
        "experience": {"years": "5–10 років", "field": "операційне управління / офіс-менеджмент, бажано в ОАЕ", "req": "pref"},
        "education": {"type": "", "req": "pref"},
        "notes": "Розширений операційний менеджер Dimax Group у Дубаї. Арабська — великий плюс навіть на розмовному рівні, але не блокер. " + CIS_NOTE,
        "prompt": "ПОСАДА: Country Manager (Dubai) у Dimax Group. Розширений операційний менеджер: організація офісу з нуля, адміністративні питання, робота з юристами, аутсорсинговими й аутстафінговими компаніями, підписання договорів.\n\n"
                  "ОБОВ'ЯЗКОВІ ВИМОГИ (якщо хоч одна не виконана — бал не вище 49):\n1. Зараз проживає в Дубаї або сусідньому еміраті ОАЕ.\n2. Англійська — вільно.\n3. Українська або російська — вільно (див. правило про походження).\n4. Досвід організації або управління офісом, адміністративних і операційних процесів.\n5. Досвід роботи з юристами, підрядниками (аутсорс/аутстаф) і договорами.\n\n"
                  "БАЖАНІ ВИМОГИ (додають бали):\n- Арабська, навіть базова розмовна (великий плюс).\n- Знання реєстрації компаній, віз і ліцензій в ОАЕ.\n- 5–10 років операційного досвіду, особливо в ОАЕ.\n- Досвід відкриття представництва або філії з нуля.\n\n"
                  "ВАГИ (разом 100): локація ОАЕ — 20; мови — 25; операційний та офісний досвід — 30; юристи, підрядники, договори — 15; арабська та знання ринку ОАЕ — 10.\n\n"
                  "ОСОБЛИВІ ПРАВИЛА:\n" + CIS_RULES + "\n\n"
                  "ШКАЛА: 85–100 — ідеальний збіг; 70–84 — сильний кандидат (наприклад, сильний операційний менеджер в ОАЕ з СНД без арабської); 50–69 — потенційний, є прогалини; 0–49 — не виконано обов'язкові вимоги.",
    },
}


def from_template(name):
    v = blank_vacancy()
    v.update(json.loads(json.dumps(TEMPLATES[name])))
    v["promptFor"] = crit_hash(v)
    return v


# =====================================================================
# Сховище (файли в робочій папці)
# =====================================================================
class Store:
    def __init__(self, work_dir):
        self.dir = work_dir
        os.makedirs(os.path.join(work_dir, "texts"), exist_ok=True)
        os.makedirs(os.path.join(work_dir, "files"), exist_ok=True)
        self.p_vac = os.path.join(work_dir, "vacancies.json")
        self.p_cand = os.path.join(work_dir, "candidates.json")
        self.p_res = os.path.join(work_dir, "results.jsonl")
        self.reload()

    def reload(self):
        self.vacancies = self._load_json(self.p_vac, {})
        self.candidates = self._load_json(self.p_cand, {})
        self.results = {}
        if os.path.exists(self.p_res):
            with open(self.p_res, encoding="utf-8") as f:
                for line in f:
                    try:
                        rec = json.loads(line)
                    except Exception:
                        continue
                    self._merge_result(rec)
        if not self.vacancies:  # перший запуск: дві вакансії Dimax
            for name in TEMPLATES:
                v = from_template(name)
                v["id"] = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")
                self.vacancies[v["id"]] = v
            self.save_vacancies()

    @staticmethod
    def _load_json(path, default):
        try:
            with open(path, encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return default

    def _write_json(self, path, data):
        os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
        tmp = path + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False)
        os.replace(tmp, path)

    def save_vacancies(self):
        self._write_json(self.p_vac, self.vacancies)

    def save_candidates(self):
        self._write_json(self.p_cand, self.candidates)

    def _merge_result(self, rec):
        key = rec["key"]
        cur = self.results.setdefault(key, {"profile": {}, "results": {}, "error": "", "status": ""})
        if rec.get("reset"):
            cur.update({"profile": {}, "results": {}, "error": "", "status": ""})
            return
        if rec.get("profile"):
            cur["profile"] = rec["profile"]
        cur["results"].update(rec.get("results") or {})
        for k in ("error", "status"):
            if k in rec:
                cur[k] = rec[k]

    def append_result(self, rec):
        self._merge_result(rec)
        with open(self.p_res, "a", encoding="utf-8") as f:
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")

    def text_path(self, key):
        return os.path.join(self.dir, "texts", stable_hash(key) + ".txt")

    def get_text(self, key):
        try:
            with open(self.text_path(key), encoding="utf-8") as f:
                return f.read()
        except Exception:
            return None

    def put_text(self, key, text):
        with open(self.text_path(key), "w", encoding="utf-8") as f:
            f.write(text)

    def clear_candidates(self):
        self.candidates, self.results = {}, {}
        self.save_candidates()
        for p in (self.p_res,):
            if os.path.exists(p):
                os.remove(p)
        for sub in ("texts", "files"):
            d = os.path.join(self.dir, sub)
            for fn in os.listdir(d):
                os.remove(os.path.join(d, fn))


# =====================================================================
# Імпорт кандидатів: CSV з Betterteam і файли резюме
# =====================================================================
def _norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def import_csvs(store: Store, sources):
    """sources: [(ім'я_файлу, файл)]. Повертає (рядків, нових, оновлених)."""
    rows = new = upd = 0
    for name, f in sources:
        df = pd.read_csv(f, dtype=str, keep_default_na=False)
        base = re.sub(r"\.csv$", "", os.path.basename(name), flags=re.I)
        tag = re.sub(r"^betterteam-[^-]+-(new-candidates-)?|-\d{8}-\d{6}$", "", base, flags=re.I) or base
        for r in df.to_dict("records"):
            rows += 1
            email = str(r.get("Candidate Email", "")).strip().lower()
            url = str(r.get("Candidate Resume URL", "")).strip()
            key = email or ("url:" + url if url else "name:" + _norm(r.get("Candidate Name", "")) + str(r.get("Candidate Phone", "")))
            c = store.candidates.get(key)
            if c is None:
                new += 1
                c = {"key": key, "name": str(r.get("Candidate Name", "")).strip(), "email": email,
                     "phone": str(r.get("Candidate Phone", "")).strip(), "resume_url": url,
                     "record_url": str(r.get("Candidate Record URL", "")).strip(),
                     "applied": str(r.get("Candidate Application Date", "")).strip(),
                     "source": str(r.get("Candidate Source", "")).strip(), "vacancy_files": [tag], "file": ""}
                store.candidates[key] = c
            else:
                upd += 1
                if tag not in c["vacancy_files"]:
                    c["vacancy_files"].append(tag)
                if url and not c.get("resume_url"):
                    c["resume_url"] = url
    store.save_candidates()
    return rows, new, upd


def _iter_upload(name, data):
    if name.lower().endswith(".zip"):
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            for info in z.infolist():
                base = os.path.basename(info.filename)
                if info.is_dir() or base.startswith(".") or "__MACOSX" in info.filename:
                    continue
                if base.lower().rsplit(".", 1)[-1] in ("pdf", "docx", "doc", "txt"):
                    yield base, z.read(info)
    elif name.lower().rsplit(".", 1)[-1] in ("pdf", "docx", "doc", "txt"):
        yield os.path.basename(name), data


def import_files(store: Store, uploads):
    """uploads: [(ім'я, bytes)] — PDF/DOCX/TXT або ZIP. Прив'язує файли до кандидатів з CSV
    (за іменем файлу з посилання або за ім'ям кандидата); решта стають окремими кандидатами."""
    by_url = {os.path.basename(c["resume_url"].split("?")[0]).lower(): k for k, c in store.candidates.items() if c.get("resume_url")}
    by_name = {}
    for k, c in store.candidates.items():
        toks = _norm(c.get("name", "")).split()
        if len(toks) >= 2:
            by_name[k] = toks
    matched = standalone = 0
    for up_name, data in uploads:
        for base, raw in _iter_upload(up_name, data):
            path = os.path.join(store.dir, "files", base)
            with open(path, "wb") as f:
                f.write(raw)
            key = by_url.get(base.lower())
            if key is None:
                nb = set(_norm(os.path.splitext(base)[0]).split())
                hits = [k for k, toks in by_name.items() if all(t in nb for t in toks)]
                key = hits[0] if len(hits) == 1 else None
            if key:
                matched += 1
            else:
                standalone += 1
                key = "file:" + base.lower()
                if key not in store.candidates:
                    nm = re.sub(r"^\d+[_\-\s]*", "", os.path.splitext(base)[0])
                    store.candidates[key] = {"key": key, "name": re.sub(r"[_\-]+", " ", nm).strip(), "email": "", "phone": "",
                                             "resume_url": "", "record_url": "", "applied": "", "source": "файл",
                                             "vacancy_files": ["завантажені файли"], "file": ""}
            store.candidates[key]["file"] = base
            if os.path.exists(store.text_path(key)):
                os.remove(store.text_path(key))  # новий файл → перечитати
            store.append_result({"key": key, "status": "", "error": ""})
    store.save_candidates()
    return matched, standalone


# =====================================================================
# Читання резюме
# =====================================================================
_ocr_engine, _ocr_checked = None, False
_ocr_lock = threading.Lock()


def _ocr():
    global _ocr_engine, _ocr_checked
    with _ocr_lock:
        if not _ocr_checked:
            _ocr_checked = True
            try:
                from rapidocr_onnxruntime import RapidOCR
                _ocr_engine = RapidOCR()
            except Exception:
                _ocr_engine = None
    return _ocr_engine


def _ocr_pdf(doc):
    engine = _ocr()
    if engine is None:
        return ""
    lines = []
    for i in range(min(3, len(doc))):
        img = doc[i].get_pixmap(dpi=150).tobytes("png")
        with _ocr_lock:
            result, _ = engine(img)
        if result:
            lines.append("\n".join(item[1] for item in result))
    return "\n".join(lines)


def extract_bytes(raw, name):
    """Повертає (текст, ознака). ознака: 'text' | 'ocr' | 'error: …'."""
    import docx
    import pymupdf as fitz
    try:
        if raw[:4] == b"%PDF" or name.lower().endswith(".pdf"):
            doc = fitz.open(stream=io.BytesIO(raw), filetype="pdf")
            text = "\n".join(p.get_text("text") for p in doc).strip()
            kind = "text"
            if len(re.sub(r"\s", "", text)) < 80:
                text, kind = _ocr_pdf(doc).strip(), "ocr"
            doc.close()
            return text, kind
        if raw[:2] == b"PK" or name.lower().endswith(".docx"):
            d = docx.Document(io.BytesIO(raw))
            parts = [p.text for p in d.paragraphs if p.text.strip()]
            for t in d.tables:
                for r in t.rows:
                    rt = " | ".join(c.text.strip() for c in r.cells if c.text.strip())
                    if rt:
                        parts.append(rt)
            return "\n".join(parts).strip(), "text"
        if raw[:4] == b"\xd0\xcf\x11\xe0" or name.lower().endswith(".doc"):
            chunks = re.findall(rb"[A-Za-z0-9\s\.,;:\-\+@/()]{5,}", raw)
            return " ".join(c.decode("latin1", errors="ignore") for c in chunks).strip(), "text"
        if name.lower().endswith(".txt"):
            return raw.decode("utf-8", errors="ignore"), "text"
        return "", "error: непідтримуваний формат"
    except Exception as e:
        return "", f"error: {e}"


def get_resume_text(store: Store, c):
    """Текст резюме кандидата: кеш → завантажений файл → посилання. Повертає (текст, помилка)."""
    t = store.get_text(c["key"])
    if t:
        return t, ""
    raw, name = None, ""
    if c.get("file"):
        p = os.path.join(store.dir, "files", c["file"])
        if os.path.exists(p):
            with open(p, "rb") as f:
                raw, name = f.read(), c["file"]
    if raw is None and c.get("resume_url"):
        try:
            r = requests.get(c["resume_url"], headers={"User-Agent": "Mozilla/5.0"}, timeout=25)
            r.raise_for_status()
            raw, name = r.content, c["resume_url"].split("?")[0]
        except Exception as e:
            code = getattr(getattr(e, "response", None), "status_code", None)
            if code in (403, 404):
                return "", f"посилання не відкривається ({code}) — завантажте файл резюме вручну"
            return "", f"не вдалося завантажити резюме: {str(e)[:120]}"
    if raw is None:
        return "", "немає ні посилання, ні файлу резюме"
    text, kind = extract_bytes(raw, name)
    if kind.startswith("error"):
        return "", kind[7:]
    text = compact(text)
    if len(re.sub(r"\s", "", text)) < 40:
        return "", "скан без тексту (OCR не допоміг)"
    store.put_text(c["key"], text)
    return text, ""


def compact(text):
    seen, out = set(), []
    for line in text.splitlines():
        line = re.sub(r"\s+", " ", line).strip()
        if len(line) < 2 or not re.search(r"\w", line) or line.lower() in seen:
            continue
        seen.add(line.lower())
        out.append(line)
    text = "\n".join(out)
    if len(text) > MAX_CV_CHARS:
        head = int(MAX_CV_CHARS * 0.75)
        text = text[:head] + "\n[...]\n" + text[-(MAX_CV_CHARS - head):]
    return text


# --- економний фільтр (для вакансій з обов'язковою укр/рос)
_CIS_WORDS = [r"ukrain\w*", "kyiv", "kiev", "kharkiv", r"odess?a", "lviv", r"dnipro\w*", r"zaporizh\w*", "vinnytsia",
              r"azerbaijan\w*", "azeri", "baku", "sumqayit", "ganja", "socar", "azercell", r"georgia\w*", "tbilisi", "batumi",
              "kutaisi", r"kazakh\w*", "almaty", "astana", "shymkent", "karaganda", "kimep", "kbtu", r"uzbek\w*", "tashkent",
              "samarkand", "bukhara", r"armenia\w*", "yerevan", "gyumri", r"moldov\w*", "chisinau", r"kyrgyz\w*", "bishkek",
              r"tajik\w*", "dushanbe", r"turkmen\w*", "ashgabat", r"belarus\w*", "minsk", "gomel", r"latvia\w*", "riga",
              r"lithuania\w*", "vilnius", "kaunas", r"estonia\w*", "tallinn", "tartu", r"russian?", "ussr", "soviet", "moscow"]
CIS_RE = re.compile(r"\b(" + "|".join(_CIS_WORDS) + r")\b", re.I)
CODES = "380|994|995|7|998|374|373|996|992|993|375|371|370|372"
PHONE_TXT = re.compile(r"\+\s?(" + CODES + r")\b")
PHONE_FLD = re.compile(r"^(" + CODES + r")")
EMAIL_TLD = re.compile(r"\.(ua|az|ge|kz|uz|am|md|kg|tj|tm|by|lv|lt|ee|ru)$", re.I)
CYR = re.compile(r"[а-яіїєґё]", re.I)
SURNAME = re.compile(r"\w+(enko|iuk|yuk|chuk|shvili|adze|dze|yan|ian|ov|ova|ev|eva|yev|yeva|zade|bayli)\b|\w{4,}(li|ly)\b", re.I)


def has_cis_signal(text, c):
    digits = re.sub(r"\D", "", c.get("phone", ""))
    return bool(CIS_RE.search(text) or PHONE_TXT.search(text) or CYR.search(text)
                or (len(digits) >= 10 and PHONE_FLD.match(digits) and not digits.startswith("971"))
                or EMAIL_TLD.search(c.get("email", "")) or SURNAME.search(c.get("name", "")))


# =====================================================================
# ШІ
# =====================================================================
class Lang(BaseModel):
    lang: str
    level: str


class Profile(BaseModel):
    name: str
    location: str = Field(description="місто, країна, де живе зараз")
    languages: list[Lang]
    experience_years: float = Field(description="загальний стаж у роках; -1 якщо невідомо")
    education: str
    education_type: Literal["технічна", "гуманітарна", "економічна", "інша", "не вказано"]
    tools: list[str]


class VacResult(BaseModel):
    vacancy_id: str
    score: int
    matched: list[str]
    missing: list[str]
    summary: str


class ScoreOut(BaseModel):
    candidate: Profile
    results: list[VacResult]


class QuotaExhausted(Exception):
    pass


def _is_quota(m):
    return any(k in m for k in ("429", "resource_exhausted", "quota", "rate limit", "rate_limit"))


def _is_hard_quota(m):
    return any(k in m for k in ("per day", "perday", "per_day", "daily", "insufficient_quota", "billing"))


class LLM:
    def __init__(self, s: Settings):
        from openai import OpenAI
        kw = {"api_key": s.api_key, "max_retries": 0, "timeout": 120}
        if s.provider == "gemini":
            kw["base_url"] = GEMINI_URL
        self.client = OpenAI(**kw)
        self.s = s
        self.reasoning = ["none", "minimal", None] if s.provider == "gemini" else [None]
        self._lock = threading.Lock()
        self._next = 0.0
        self.quota_over = threading.Event()
        self.calls = 0
        self.tokens_in = self.tokens_out = 0

    def _wait(self):
        pause = self.s.request_pause
        if pause <= 0:
            return
        with self._lock:
            now = time.monotonic()
            slot = max(now, self._next)
            self._next = slot + pause
        time.sleep(max(0.0, slot - now))

    def call(self, messages, schema=None, max_tokens=1200, think=False):
        free = self.s.mode == "free"
        quota_hits = transient = 0
        for _ in range(8):
            if self.quota_over.is_set():
                raise QuotaExhausted()
            self._wait()
            kw = dict(model=self.s.model_name, messages=messages, temperature=0.2, max_tokens=max_tokens)
            effort = None if think else self.reasoning[0]
            if effort:
                kw["reasoning_effort"] = effort
            try:
                if schema:
                    r = self.client.beta.chat.completions.parse(response_format=schema, **kw)
                else:
                    r = self.client.chat.completions.create(**kw)
                with self._lock:
                    self.calls += 1
                    u = getattr(r, "usage", None)
                    if u:
                        self.tokens_in += getattr(u, "prompt_tokens", 0) or 0
                        self.tokens_out += getattr(u, "completion_tokens", 0) or 0
                msg = r.choices[0].message
                if schema:
                    if msg.parsed is None:
                        raise ValueError("порожня відповідь моделі")
                    return msg.parsed
                return msg.content or ""
            except QuotaExhausted:
                raise
            except Exception as e:
                m = str(e).lower()
                if "reasoning" in m and effort is not None:
                    with self._lock:
                        if self.reasoning and self.reasoning[0] == effort:
                            self.reasoning.pop(0)
                    continue
                if _is_quota(m):
                    quota_hits += 1
                    if _is_hard_quota(m) or quota_hits >= (3 if free else 4):
                        self.quota_over.set()
                        raise QuotaExhausted() from e
                    time.sleep(60 if free else 20)
                    continue
                if transient < 2 and any(k in m for k in ("timeout", "timed out", "500", "502", "503", "unavailable", "connection")):
                    transient += 1
                    time.sleep(10)
                    continue
                raise
        raise RuntimeError("забагато повторних спроб")


def generator_prompt(v):
    return f"""Ти — досвідчений рекрутер і автор інструкцій для ШІ. На основі критеріїв вакансії склади промпт, за яким інший ШІ оцінюватиме резюме кандидатів за шкалою 0–100.

Вимоги до промпту:
- Українською, коротко й конкретно, без вступу та пояснень від себе. Тільки сам промпт, без Markdown-розмітки (без **, #).
- Структура з такими розділами великими літерами:
ПОСАДА: одне-два речення про роль.
ОБОВ'ЯЗКОВІ ВИМОГИ (якщо хоч одна не виконана — бал не вище 49): нумерований перелік.
БАЖАНІ ВИМОГИ (додають бали): перелік.
ВАГИ (разом 100): розподіл балів між групами критеріїв.
ОСОБЛИВІ ПРАВИЛА: як трактувати непрямі ознаки, винятки, що вважати виконаним.
ШКАЛА: що означають діапазони 85–100, 70–84, 50–69, 0–49.
- Кожну вимогу сформулюй так, щоб її можна було перевірити за текстом резюме. Розшифруй загальні поняття (наприклад, «технічна освіта» = інженерія, IT, фізика, енергетика тощо; «вільна англійська» = C1+, робота або навчання англійською).
- Не вигадуй вимог, яких немає в критеріях. «обов'язково» → обов'язкові вимоги; «бажано» → бажані.

КРИТЕРІЇ ВАКАНСІЇ:
{criteria_text(v)}"""


def generate_prompt(s: Settings, v):
    text = LLM(s).call([{"role": "user", "content": generator_prompt(v)}], max_tokens=2500, think=True)
    return text.replace("**", "").strip()


def scoring_system(vacs):
    # Незмінна частина йде першою: Gemini кешує однаковий початок запиту і бере за нього менше.
    return ("Ти — ШІ-рекрутер. Оціни кандидата окремо для кожної вакансії нижче, суворо за її промптом. Бали — цілі від 0 до 100.\n"
            "У відповіді: candidate — профіль з резюме (мови й навички українською; experience_years = -1, якщо стаж невідомий); "
            f"results — рівно {len(vacs)} елемент(и), по одному на кожен vacancy_id; matched — коротко, які вимоги виконано; "
            "missing — чого бракує або що перевірити; summary — одне речення українською до 30 слів.\n\n=== ВАКАНСІЇ ===\n"
            + "\n\n".join(f"--- vacancy_id: {v['id']} | {v['title']}\n{v['prompt']}" for v in vacs))


def score_candidate(llm: LLM, vacs, c, text):
    user = f"Кандидат: {c.get('name', '')}\nТелефон: {c.get('phone', '')}\n\n=== РЕЗЮМЕ ===\n{text}"
    out = llm.call([{"role": "system", "content": scoring_system(vacs)}, {"role": "user", "content": user}],
                   schema=ScoreOut, max_tokens=300 + 350 * len(vacs))
    res = {}
    for v in vacs:
        r = next((x for x in out.results if x.vacancy_id == v["id"]), None)
        if r is None and len(vacs) == 1 and out.results:
            r = out.results[0]
        if r is None:
            continue
        res[v["id"]] = {"score": max(0, min(100, int(r.score))), "matched": r.matched[:10], "missing": r.missing[:10],
                        "summary": r.summary, "promptHash": stable_hash(v["prompt"]), "at": int(time.time())}
    if not res:
        raise ValueError("ШІ не повернув оцінок")
    p = out.candidate.model_dump()
    if p.get("experience_years", -1) < 0:
        p["experience_years"] = None
    return p, res


# =====================================================================
# Запуск оцінки
# =====================================================================
def needs_scoring(store: Store, key, vacs, rescore=False):
    r = store.results.get(key, {})
    if r.get("status") == "unreadable" and not rescore:
        return False
    if rescore or r.get("error"):
        return True
    return any(v["id"] not in r.get("results", {}) or r["results"][v["id"]].get("promptHash") != stable_hash(v["prompt"]) for v in vacs)


def process_one(store: Store, llm: LLM, s: Settings, key, vacs, rescore):
    c = store.candidates[key]
    try:
        if llm.quota_over.is_set():
            return None
        text, err = get_resume_text(store, c)
        if err:
            return {"key": key, "status": "unreadable", "error": err}
        cur = store.results.get(key, {}).get("results", {})
        todo = [v for v in vacs if rescore or v["id"] not in cur or cur[v["id"]].get("promptHash") != stable_hash(v["prompt"])]
        rec = {"key": key, "status": "ok", "error": "", "results": {}}
        if s.smart_filter and not has_cis_signal(text, c):
            for v in [v for v in todo if requires_uk_ru(v)]:
                rec["results"][v["id"]] = {"score": 5, "matched": [], "missing": ["немає ознак української/російської мови чи походження з UA/СНД/Балтії"],
                                           "summary": "Відсіяно економним фільтром без запиту до ШІ.", "promptHash": stable_hash(v["prompt"]),
                                           "at": int(time.time()), "filtered": True}
            todo = [v for v in todo if not requires_uk_ru(v)]
        if todo:
            profile, res = score_candidate(llm, todo, c, text)
            rec["profile"] = profile
            rec["results"].update(res)
            if profile.get("name") and len(profile["name"]) > 2 and c.get("source") == "файл":
                rec["name"] = profile["name"]
        return rec
    except QuotaExhausted:
        return None
    except Exception as e:
        return {"key": key, "status": "error", "error": f"Помилка ШІ: {str(e)[:200]}"}


def run_scoring(store: Store, s: Settings, keys, vacs, rescore=False, on_progress: Optional[Callable] = None):
    llm = LLM(s)
    st = {"total": len(keys), "done": 0, "ai": 0, "filtered": 0, "unreadable": 0, "errors": 0, "started": time.time()}
    it = iter(keys)
    inflight = set()
    w = s.n_workers
    with concurrent.futures.ThreadPoolExecutor(max_workers=w) as ex:
        def refill():
            while len(inflight) < w * 2 and not llm.quota_over.is_set():
                try:
                    k = next(it)
                except StopIteration:
                    return
                inflight.add(ex.submit(process_one, store, llm, s, k, vacs, rescore))
        refill()
        while inflight:
            done, _ = concurrent.futures.wait(inflight, return_when=concurrent.futures.FIRST_COMPLETED)
            for f in done:
                inflight.discard(f)
                rec = f.result()
                if rec is None:
                    continue
                name = rec.pop("name", None)
                if name:
                    store.candidates[rec["key"]]["name"] = name
                store.append_result(rec)
                st["done"] += 1
                if rec["status"] == "unreadable":
                    st["unreadable"] += 1
                elif rec["status"] == "error":
                    st["errors"] += 1
                elif rec.get("profile"):
                    st["ai"] += 1
                else:
                    st["filtered"] += 1
                if on_progress:
                    on_progress(st, store.candidates[rec["key"]])
            refill()
    store.save_candidates()
    st.update(quota_exhausted=llm.quota_over.is_set(), remaining=st["total"] - st["done"], seconds=time.time() - st["started"],
              calls=llm.calls, tokens_in=llm.tokens_in, tokens_out=llm.tokens_out)
    return st


# =====================================================================
# Таблиця результатів і Excel
# =====================================================================
def verdict(s):
    return "Підходить" if s >= 70 else "Можливо" if s >= 50 else "Не підходить"


def lang_str(p):
    return ", ".join(l["lang"] + (f" ({l['level']})" if l.get("level") and l["level"] != "не вказано" else "") for l in (p or {}).get("languages", []))


def results_frame(store: Store):
    rows = []
    vacs = list(store.vacancies.values())
    for k, c in store.candidates.items():
        r = store.results.get(k, {})
        p = r.get("profile") or {}
        row = {"key": k, "Кандидат": c.get("name", ""), "Статус": {"ok": "оцінено", "unreadable": "не прочитано", "error": "помилка ШІ"}.get(r.get("status"), "не оцінено"),
               "Проблема": r.get("error", ""), "Мови": lang_str(p), "Локація": p.get("location", ""),
               "Досвід, р.": p.get("experience_years"), "Освіта": p.get("education", ""), "Тип освіти": p.get("education_type", ""),
               "Навички": ", ".join(p.get("tools", [])), "Email": c.get("email", ""), "Телефон": c.get("phone", ""),
               "Резюме": c.get("resume_url", "") or c.get("file", ""), "Betterteam": c.get("record_url", ""),
               "Вакансії відгуку": " + ".join(c.get("vacancy_files", [])), "Дата відгуку": c.get("applied", "")}
        for v in vacs:
            x = r.get("results", {}).get(v["id"])
            row[f"bal::{v['id']}"] = x["score"] if x else None
            row[f"sum::{v['id']}"] = x["summary"] if x else ""
            row[f"stale::{v['id']}"] = bool(x and x.get("promptHash") != stable_hash(v.get("prompt", "")))
        rows.append(row)
    return pd.DataFrame(rows)


def build_excel(store: Store, threshold=55):
    vacs = list(store.vacancies.values())
    buf = io.BytesIO()
    legend = [("Аркуш / колонка", "Пояснення")]
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        sheets = []
        for v in vacs:
            rows = []
            for k, c in store.candidates.items():
                r = store.results.get(k, {})
                x = r.get("results", {}).get(v["id"])
                if not x:
                    continue
                p = r.get("profile") or {}
                rows.append({"Кандидат": c.get("name", ""), "Бал": x["score"], "Висновок": verdict(x["score"]),
                             "Що збіглося": "; ".join(x.get("matched", [])), "Чого бракує / перевірити": "; ".join(x.get("missing", [])),
                             "Коротко": x.get("summary", ""), "Мови": lang_str(p), "Локація": p.get("location", ""),
                             "Досвід, р.": p.get("experience_years"), "Освіта": p.get("education", ""), "Тип освіти": p.get("education_type", ""),
                             "Навички": ", ".join(p.get("tools", [])), "Email": c.get("email", ""), "Телефон": c.get("phone", ""),
                             "Резюме": c.get("resume_url", "") or c.get("file", ""), "Betterteam": c.get("record_url", ""),
                             "Вакансії відгуку": " + ".join(c.get("vacancy_files", [])), "Дата відгуку": c.get("applied", "")})
            df = pd.DataFrame(rows)
            if df.empty:
                continue
            df = df.sort_values("Бал", ascending=False)
            name = re.sub(r"[\[\]:*?/\\]", " ", v["title"])[:28].strip() or "Вакансія"
            while name in sheets:
                name = name[:25] + "_2"
            sheets.append(name)
            df.to_excel(w, sheet_name=name, index=False)
            legend.append((name, f"Усі оцінені кандидати для «{v['title']}», від найвищого балу. Зелений — 70+, жовтий — {threshold}–69."))
        att = results_frame(store)
        if not att.empty:
            att = att[att["Статус"] != "оцінено"][["Кандидат", "Статус", "Проблема", "Email", "Телефон", "Резюме", "Betterteam"]]
            att.to_excel(w, sheet_name="Потребує уваги", index=False)
            legend.append(("Потребує уваги", "Резюме, які не вдалося прочитати чи оцінити, або ще не оцінені. Причина — у колонці «Проблема»."))
        pd.DataFrame([{"Вакансія": v["title"], "Критерії": criteria_text(v), "Промпт оцінки": v.get("prompt", "")} for v in vacs]).to_excel(w, sheet_name="Промпти", index=False)
        legend += [("Промпти", "Критерії та промпт, за якими оцінювали кандидатів."),
                   ("Важливо", "Оцінки зроблені ШІ і є орієнтиром для відбору, а не остаточним рішенням."),
                   ("Сформовано", datetime.now().strftime("%d.%m.%Y %H:%M"))]
        pd.DataFrame(legend[1:], columns=legend[0]).to_excel(w, sheet_name="Легенда", index=False)
    wb = load_workbook(io.BytesIO(buf.getvalue()))
    wb.move_sheet("Легенда", offset=-len(wb.sheetnames) + 1)
    head = PatternFill("solid", fgColor="1F3A68")
    green, yellow = PatternFill("solid", fgColor="C6EFCE"), PatternFill("solid", fgColor="FFEB9C")
    widths = {"Кандидат": 24, "Що збіглося": 45, "Чого бракує / перевірити": 45, "Коротко": 60, "Мови": 32, "Освіта": 30,
              "Навички": 36, "Критерії": 70, "Промпт оцінки": 110, "Проблема": 50, "Пояснення": 100, "Аркуш / колонка": 26,
              "Email": 26, "Резюме": 12, "Betterteam": 12}
    for ws in wb.worksheets:
        hdr = [c.value for c in ws[1]]
        for c in ws[1]:
            c.font = Font(name="Arial", bold=True, color="FFFFFF")
            c.fill = head
            c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.freeze_panes = "B2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
        for j, h in enumerate(hdr, 1):
            ws.column_dimensions[get_column_letter(j)].width = widths.get(h, 14)
        col = {h: j for j, h in enumerate(hdr, 1)}
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial")
                if hdr[cell.column - 1] in ("Що збіглося", "Чого бракує / перевірити", "Коротко", "Критерії", "Промпт оцінки", "Пояснення", "Проблема"):
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            if "Бал" in col:
                cell = row[col["Бал"] - 1]
                if isinstance(cell.value, (int, float)):
                    cell.fill = green if cell.value >= 70 else yellow if cell.value >= threshold else PatternFill()
            for h in ("Резюме", "Betterteam"):
                if h in col:
                    cell = row[col[h] - 1]
                    if isinstance(cell.value, str) and cell.value.startswith("http"):
                        cell.hyperlink = cell.value
                        cell.value = "Відкрити"
                        cell.style = "Hyperlink"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# =====================================================================
# Резервна копія всіх даних (для хостингів, що стирають файли при перезапуску)
# =====================================================================
def backup_zip(store: Store):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for root, _, files in os.walk(store.dir):
            for fn in files:
                if fn.endswith(".tmp"):
                    continue
                full = os.path.join(root, fn)
                z.write(full, os.path.relpath(full, store.dir))
    return buf.getvalue()


def restore_zip(store: Store, data):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = z.namelist()
        if "vacancies.json" not in names:
            raise ValueError("це не резервна копія Dimax HR (немає vacancies.json)")
        for n in names:
            if n.startswith("/") or ".." in n.split("/"):
                continue
            target = os.path.join(store.dir, n)
            os.makedirs(os.path.dirname(target), exist_ok=True)
            with open(target, "wb") as f:
                f.write(z.read(n))
    store.reload()
